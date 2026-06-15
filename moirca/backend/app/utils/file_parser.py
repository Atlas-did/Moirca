"""
文件解析工具（继承自MiroFish）
支持PDF、Markdown、TXT文件的文本提取
"""
import os
from pathlib import Path
from typing import List


def _read_text_with_fallback(file_path: str) -> str:
    data = Path(file_path).read_bytes()
    try:
        return data.decode('utf-8')
    except UnicodeDecodeError:
        pass
    encoding = None
    try:
        from charset_normalizer import from_bytes
        best = from_bytes(data).best()
        if best and best.encoding:
            encoding = best.encoding
    except Exception:
        pass
    if not encoding:
        try:
            import chardet
            result = chardet.detect(data)
            encoding = result.get('encoding') if result else None
        except Exception:
            pass
    if not encoding:
        encoding = 'utf-8'
    return data.decode(encoding, errors='replace')


class FileParser:
    SUPPORTED_EXTENSIONS = {'.pdf', '.docx', '.md', '.markdown', '.txt', '.json', '.csv'}

    @classmethod
    def is_supported(cls, file_path: str) -> bool:
        return Path(file_path).suffix.lower() in cls.SUPPORTED_EXTENSIONS

    @classmethod
    def extract_text(cls, file_path: str) -> str:
        path = Path(file_path)
        if not path.exists():
            raise FileNotFoundError(f"文件不存在: {file_path}")
        suffix = path.suffix.lower()
        if suffix == '.pdf':
            return cls._extract_from_pdf(file_path)
        elif suffix == '.docx':
            return cls._extract_from_docx(file_path)
        elif suffix in {'.md', '.markdown', '.txt', '.json', '.csv'}:
            return _read_text_with_fallback(file_path)
        raise ValueError(f"不支持的文件格式: {suffix}")

    @staticmethod
    def _extract_from_pdf(file_path: str) -> str:
        try:
            import fitz
        except ImportError:
            raise ImportError("需要安装PyMuPDF: pip install PyMuPDF")
        text_parts = []
        with fitz.open(file_path) as doc:
            for page in doc:
                text = page.get_text()
                if text.strip():
                    text_parts.append(text)
        return "\n\n".join(text_parts)

    @staticmethod
    def _extract_from_docx(file_path: str) -> str:
        try:
            from docx import Document
        except ImportError:
            raise ImportError("需要安装 python-docx: pip install python-docx")
        doc = Document(file_path)
        parts = []
        for p in doc.paragraphs:
            t = (p.text or "").strip()
            if t:
                parts.append(t)
        # 表格内容也提取一下（常见于计划书/模板）
        for table in doc.tables:
            for row in table.rows:
                cells = [((c.text or "").strip()) for c in row.cells]
                line = "\t".join([c for c in cells if c])
                if line.strip():
                    parts.append(line)
        return "\n".join(parts).strip()


def split_text_into_chunks(text: str, chunk_size: int = 500, overlap: int = 50) -> List[str]:
    if len(text) <= chunk_size:
        return [text] if text.strip() else []
    chunks = []
    start = 0
    while start < len(text):
        end = start + chunk_size
        if end < len(text):
            for sep in ['。', '！', '？', '\n\n', '. ', '! ', '? ']:
                last_sep = text[start:end].rfind(sep)
                if last_sep != -1 and last_sep > chunk_size * 0.3:
                    end = start + last_sep + len(sep)
                    break
        chunk = text[start:end].strip()
        if chunk:
            chunks.append(chunk)
        start = end - overlap if end < len(text) else len(text)
    return chunks
